"""Generate 3 trap-filled test workbooks in the exact template format (row1 title, row2 headers, row3 example, data from row 4)."""
import openpyxl, shutil, os
SRC = "/workspace/template.xlsx"
OUT = os.path.join(os.path.dirname(__file__), "..", "tests")
PUB = os.path.join(os.path.dirname(__file__), "..", "public")

def build(name, apps, ters, keep_stores=True, keep_example=True):
    wb = openpyxl.load_workbook(SRC)
    wa, wt = wb["Applications"], wb["Territories"]
    if not keep_example:
        for c in range(1, 15): wa.cell(3, c).value = None
        for c in range(1, 8): wt.cell(3, c).value = None
    for i, r in enumerate(apps):
        for j, v in enumerate(r): wa.cell(4 + i, 1 + j).value = v
    for i, r in enumerate(ters):
        for j, v in enumerate(r): wt.cell(4 + i, 1 + j).value = v
    if not keep_stores: del wb["Existing stores"]
    p = os.path.join(OUT, name); wb.save(p); shutil.copy(p, os.path.join(PUB, name)); print("wrote", p)

K = "Karnataka"
# ---------- Test 1: mixed traps (keeps template example rows, which must be skipped) ----------
ters1 = [
    ["TER-B1", "Bengaluru", K, "560001, 560011, 560041, 560070", "Tier 1", 2, "High"],
    ["TER-B2", "Bengaluru", K, "560078, 560102", "Tier 1", 1, "Medium"],
    ["TER-M1", "Mysuru", K, "570001, 570002, 570008, 570009", "Tier 2", 2, "High"],
    ["TER-MG", "Mangaluru", K, "575002, 575003", "Tier 2", 1, "Low"],
    ["TER-T1", "Tumakuru", K, "572101, 572102", "Tier 3", 1, "Medium"],
    ["TER-D1", "Davanagere", K, "577001, 577002", "Tier 3", 1, "Low"],
    ["TER-H1", "Hubballi-Dharwad", K, "580029, 580031", "Tier 2", 1, "Medium"],
]
apps1 = [
    ["APP-002", "Lakshmi Narayan", "Bengaluru", K, 560011, 60, 6, "Runs an optical shop in Jayanagar for 6 years", "Owned", 800, "Ground", 6, "Kannada", ""],
    ["APP-003", "Arjun Shetty", "bangalore", K, 560041, "45 lakh", 3, "Runs a mobile accessories shop", "Rented", 600, "Ground", 4, "English", "Can start quickly"],
    ["APP-004", "Fatima Begum", "Bengaluru", K, 560041, 55, 5, "Runs an apparel store", "Rented", 700, "Mall unit", 5, "Hindi", ""],
    ["APP-005", "Suresh Gowda", "Mysuru", K, 570002, 30, 2, "Runs a kirana store", "Rented", 500, "Ground", 3, "Kannada", "Will take a bank loan for the rest"],
    ["APP-006", "Ramesh Rao", "Mysuru", K, 570009, 50, 4, "Runs a footwear shop", "Owned", 700, "Ground", 4, "Kannada", ""],
    ["APP-007", "Priya Hegde", "Mysuru", K, 570001, 40, 0, "My father's saree shop, I help there on weekends", "Rented", 550, "Upper", 6, "Kannada", ""],
    ["APP-008", "Kiran Patil", "Mysuru", K, 570008, 45, 0, "Ran my own mobile shop for 5 years in Mandya", "Owned", 900, "Ground", 5, "English", ""],
    ["APP-009", "Anil Kumar", "Mysuru", K, 560034, 50, 4, "Runs a bakery", "Rented", 600, "Ground", 4, "Kannada", ""],
    ["APP-010", "Deepa S", "Bengaluru", K, 560099, 70, 5, "Runs a cosmetics store", "Rented", 700, "Ground", 4, "English", ""],
    ["APP-011", "Mohan Raj", "Chennai", "Tamil Nadu", 600001, 80, 6, "Runs a shoe store", "Owned", 800, "Ground", 4, "English", ""],
    ["APP-012", "Sunita Rao", "Tumakuru", K, 572101, 25, 3, "Runs a stationery shop", "Not yet identified", None, None, 4, "Kannada", "Still looking for a shop near the bus stand"],
    ["APP-013", "Vijay Kumar", "Davanagere", K, 577001, 2500000, 4, "Runs an electronics shop", "Rented", 600, "Ground", 4, "Hindi", ""],
    ["APP-002", "Lakshmi N (duplicate)", "Bengaluru", K, 560070, 60, 6, "Optical shop", "Owned", 800, "Ground", 6, "Kannada", "Submitted twice"],
    [None, "No ID Person", "Bengaluru", K, 560070, 50, 2, "Shop", "Rented", 600, "Ground", 4, "English", ""],
    ["APP-015", "Ganesh Bhat", "Mangaluru", K, 575002, 12, 1, "Teacher", "Rented", 400, "Ground", 6, "Kannada", ""],
    ["APP-016", "Rekha Desai", "Hubballi", K, 580029, 50, 7, "Running a footwear store", "Rented", None, "Ground", 5, "Kannada", ""],
    ["APP-017", "Imran Khan", "Hubballi-Dharwad", K, 580031, 60, 10, "Already run a Titan Eye+ franchise", "Owned", 700, "Ground", 4, "Hindi", "Want exclusive rights for the whole city"],
    ["APP-018", "Kavya M", "Belagavi", K, 590006, 45, 3, "Runs a gift shop", "Rented", 600, "Ground", 4, "Kannada", ""],
    ["APP-019", "Naveen Reddy", "Bengaluru", K, 560070, 42, 1, "IT engineer", "Not yet identified", 400, None, 6, "English", ""],
    ["APP-020", "Shwetha Iyer", "Bengaluru", K, 560078, 70, 8, "Ran a garments boutique", "Rented", 1200, "Ground", 2, None, ""],
    ["APP-021", "Rahul Jain", "Bengaluru", K, 560102, 41, 2, None, "Rented", 350, "Upper", 12, "English", ""],
    ["APP-022", "Meena Gowda", "Mysuru", K, "570 002", 36, "4 yrs", "Runs a pharmacy", "Rented", "650 sq ft", "Ground", 4, "Kannada", "ಅಂಗಡಿ ಬಾಡಿಗೆಗೆ ಸಿಕ್ಕಿದೆ, ದೀಪಾವಳಿಗೆ ಮುಂಚೆ ತೆರೆಯಬೇಕು"],
    ["APP-023", "Pooja Sharma", "Davanagere", K, 577002, 22, 0, "Homemaker", "Rented", 500, "Ground", 6, "Hindi", "Loan from bank pending"],
]
build("test_1_mixed_traps.xlsx", apps1, ters1)

# ---------- Test 2: dirty data, NO Existing stores sheet (built-in list must be used), example row removed ----------
ters2 = [
    ["TER-01", "Bengaluru", K, "560011,560041;560076", "Tier 1", 2, "high"],
    ["TER-02", "Mysuru", K, "570002 570008", "tier 2", "2 ", "Medium"],
    ["TER-03", "Mysuru", K, "570008, 570010", "Tier 2", 1, "Low"],  # 570008 also in TER-02 -> warning
    ["TER-04", "Tumakuru", K, "572101", "T3", 1, "Medium"],
    ["TER-05", "Belagavi", K, "590006, 590001", "Tier 2", 1, "High"],
]
apps2 = [
    ["  app-030 ", "  Raghav   Menon ", "Bangalore ", "karnataka ", "560011", "₹50 L", "6", "  RUNS A WATCH SHOP ", "owned", 750, "ground floor", 6, "english", ""],
    ["APP-031", "Divya K", "BENGALURU", K, 560041.0, "1.2 crore", 3, "Software engineer", "Rented", 600, "Ground", 8, "Kannada", ""],
    ["APP-032", "Harish", "Bengaluru", K, 560076, 50, 5, "Runs a shoe shop", "Rented", 700, "Ground", 5, "English", ""],
    ["APP-033", "Manjunath", "mysore", K, "570002", 40, 0, "It's my father's shop, I have been running it for 3 years", "Rented", 600, "Ground", 5, "Hindi", ""],
    ["APP-034", "Shilpa", "Mysuru", K, 570008, 50, 4, "Runs a boutique", "Owned", None, "Ground", 4, "Kannada", ""],
    ["APP-035", "Prakash", "Mysuru", K, 570010, 38, "five", "Runs a small grocery shop", "Rented", 500, "1st floor", 6, "Kannada", ""],
    ["APP-036", "Asha", "Tumkur", K, 572101, 21, 2, "Runs a mobile recharge shop", "Rented", 420, "Ground", 3, "", "Will open after Ugadi"],
    ["APP-037", "Nikhil", "Belgaum", K, 590006, None, 5, "Electronics shop owner", "Owned", 800, "Ground", 4, "Kannada", ""],
    ["APP-038", "Sameer", "Belagavi", K, 590001, 60, 9, "Optical store", "Owned", 900, "Ground", 4, "Hindi", ""],
    ["APP-039", "Latha", "Mysuru", K, 570002, 52, 6, "Runs a pharmacy", "Rented", 800, "Not known", 4, "Tamil", "Rented a shop but still searching for a better location"],
]
build("test_2_dirty_data.xlsx", apps2, ters2, keep_stores=False, keep_example=False)

# ---------- Test 3: slot competition, ties, proximity, quality bar ----------
ters3 = [
    ["TER-JN", "Bengaluru", K, "560011, 560041, 560069, 560070, 560004", "Tier 1", 2, "High"],
    ["TER-MY", "Mysuru", K, "570002, 570004, 570023", "Tier 2", 1, "Low"],
]
apps3 = [
    ["APP-101", "Akash", "Bengaluru", K, 560011, 80, 8, "Runs an optical store", "Owned", 800, "Ground", 6, "English", ""],
    ["APP-102", "Bhavana", "Bengaluru", K, 560041, 78, 8, "Runs an optical store", "Rented", 800, "Ground", 6, "Kannada", ""],
    ["APP-103", "Chetan", "Bengaluru", K, 560069, 75, 7, "Runs a jewellery shop", "Rented", 750, "Ground", 12, "Kannada", ""],
    ["APP-104", "Dinesh", "Bengaluru", K, 560004, 74, 7, "Runs a watch shop", "Rented", 700, "Ground", 2, "English", ""],
    ["APP-105", "Esha", "Bengaluru", K, 560070, 41, 1, "Student", "Rented", 300, "Upper", 3, "Hindi", ""],
    ["APP-106", "Farhan", "Mysuru", K, 570002, 36, 1, "Delivery driver", "Rented", 350, "Upper", 2, "Hindi", ""],
    ["APP-107", "Gita", "Mysuru", K, 570004, 70, 6, "Runs a saree shop", "Owned", 900, "Ground", 6, "Kannada", ""],
    ["APP-108", "Harsha", "Mysuru", K, 570023, 72, 7, "Runs an optical store", "Owned", 900, "Ground", 6, "Kannada", "Shop is ~1 km from your Mysuru store"],
]
build("test_3_competition.xlsx", apps3, ters3)
